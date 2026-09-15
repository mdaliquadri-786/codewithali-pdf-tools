#!/usr/bin/env python3
"""
CodeWithAli PDF Engine
High-performance engine for PDF processing using pypdf, reportlab, and Pillow.
Provides complete implementations for Merge, Split, Compress, Image-to-PDF, Watermark, and more.
"""

import sys
import os
import json
import io
import zipfile
import math
from pypdf import PdfReader, PdfWriter
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib import pagesizes, colors

def merge_pdfs(input_paths, output_path):
    writer = PdfWriter()
    for path in input_paths:
        if not os.path.exists(path):
            continue
        reader = PdfReader(path)
        for page in reader.pages:
            writer.add_page(page)
    with open(output_path, "wb") as f:
        writer.write(f)
    return {"success": True, "output": output_path, "page_count": len(writer.pages)}

def split_pdf(input_path, output_dir, mode="all", page_range=""):
    os.makedirs(output_dir, exist_ok=True)
    reader = PdfReader(input_path)
    total_pages = len(reader.pages)
    created_files = []

    if mode == "all":
        zip_filename = os.path.join(output_dir, "split_pages.zip")
        with zipfile.ZipFile(zip_filename, 'w', zipfile.ZIP_DEFLATED) as zipf:
            for idx, page in enumerate(reader.pages):
                writer = PdfWriter()
                writer.add_page(page)
                page_filename = f"page_{idx + 1}.pdf"
                page_path = os.path.join(output_dir, page_filename)
                with open(page_path, "wb") as f:
                    writer.write(f)
                zipf.write(page_path, arcname=page_filename)
                created_files.append(page_path)
        return {"success": True, "output": zip_filename, "is_zip": True, "file_count": len(created_files)}

    elif mode == "range":
        pages_to_extract = set()
        parts = [p.strip() for p in page_range.split(",") if p.strip()]
        for part in parts:
            if "-" in part:
                sub = part.split("-")
                try:
                    start = max(1, int(sub[0]))
                    end = min(total_pages, int(sub[1]))
                    for p in range(start, end + 1):
                        pages_to_extract.add(p - 1)
                except ValueError:
                    pass
            else:
                try:
                    p = int(part)
                    if 1 <= p <= total_pages:
                        pages_to_extract.add(p - 1)
                except ValueError:
                    pass

        pages_to_extract = sorted(list(pages_to_extract))
        if not pages_to_extract:
            pages_to_extract = list(range(total_pages))

        single_output = os.path.join(output_dir, "split_extracted.pdf")
        writer = PdfWriter()
        for p_idx in pages_to_extract:
            writer.add_page(reader.pages[p_idx])
        with open(single_output, "wb") as f:
            writer.write(f)
        return {"success": True, "output": single_output, "is_zip": False, "page_count": len(pages_to_extract)}

def compress_pdf(input_path, output_path, level="recommended"):
    reader = PdfReader(input_path)
    writer = PdfWriter()

    for page in reader.pages:
        writer.add_page(page)

    for page in writer.pages:
        try:
            page.compress_content_streams()
        except Exception:
            pass

    quality = 70 if level == "recommended" else (40 if level == "extreme" else 90)
    
    for page in writer.pages:
        try:
            for img in page.images:
                try:
                    img_data = img.data
                    pil_img = Image.open(io.BytesIO(img_data))
                    if pil_img.mode in ("RGBA", "P"):
                        pil_img = pil_img.convert("RGB")
                    buf = io.BytesIO()
                    pil_img.save(buf, format="JPEG", quality=quality, optimize=True)
                    img.replace(buf.getvalue())
                except Exception:
                    pass
        except Exception:
            pass

    with open(output_path, "wb") as f:
        writer.write(f)

    original_size = os.path.getsize(input_path)
    compressed_size = os.path.getsize(output_path)
    reduction = max(0, round((1 - (compressed_size / original_size)) * 100, 1)) if original_size > 0 else 0

    return {
        "success": True,
        "output": output_path,
        "original_size": original_size,
        "compressed_size": compressed_size,
        "reduction_percentage": reduction
    }

def images_to_pdf(image_paths, output_path, orientation="portrait", margin=20, page_size="a4"):
    if page_size.lower() == "letter":
        base_w, base_h = pagesizes.letter
    else:
        base_w, base_h = pagesizes.A4

    if orientation == "landscape":
        page_w, page_h = max(base_w, base_h), min(base_w, base_h)
    elif orientation == "portrait":
        page_w, page_h = min(base_w, base_h), max(base_w, base_h)
    else:
        page_w, page_h = base_w, base_h

    c = canvas.Canvas(output_path)
    margin = float(margin)

    for img_path in image_paths:
        if not os.path.exists(img_path):
            continue
        try:
            with Image.open(img_path) as img:
                img_w, img_h = img.size

                if orientation == "fit":
                    cur_w, cur_h = img_w, img_h
                    c.setPageSize((cur_w, cur_h))
                    c.drawImage(img_path, 0, 0, width=cur_w, height=cur_h)
                    c.showPage()
                else:
                    c.setPageSize((page_w, page_h))
                    avail_w = page_w - 2 * margin
                    avail_h = page_h - 2 * margin
                    scale = min(avail_w / img_w, avail_h / img_h)
                    draw_w = img_w * scale
                    draw_h = img_h * scale
                    x = (page_w - draw_w) / 2
                    y = (page_h - draw_h) / 2
                    c.drawImage(img_path, x, y, width=draw_w, height=draw_h)
                    c.showPage()
        except Exception:
            continue

    c.save()
    return {"success": True, "output": output_path}

def add_watermark(input_path, output_path, text="CONFIDENTIAL", position="diagonal", opacity=0.3, font_size=48, color_hex="#666666"):
    reader = PdfReader(input_path)
    writer = PdfWriter()

    color_hex = color_hex.lstrip("#")
    if len(color_hex) == 6:
        r = int(color_hex[0:2], 16) / 255.0
        g = int(color_hex[2:4], 16) / 255.0
        b = int(color_hex[4:6], 16) / 255.0
    else:
        r, g, b = 0.5, 0.5, 0.5

    for page in reader.pages:
        pw = float(page.mediabox.width)
        ph = float(page.mediabox.height)

        packet = io.BytesIO()
        can = canvas.Canvas(packet, pagesize=(pw, ph))
        can.setFillColor(colors.Color(r, g, b, alpha=float(opacity)))
        can.setFont("Helvetica-Bold", float(font_size))

        if position == "diagonal":
            can.saveState()
            can.translate(pw / 2, ph / 2)
            angle = math.degrees(math.atan2(ph, pw))
            can.rotate(angle)
            can.drawCentredString(0, -float(font_size) / 4, text)
            can.restoreState()
        elif position == "center":
            can.drawCentredString(pw / 2, ph / 2, text)
        elif position == "bottom-right":
            can.drawRightString(pw - 30, 30, text)
        elif position == "top-left":
            can.drawString(30, ph - 40, text)

        can.save()
        packet.seek(0)
        watermark_pdf = PdfReader(packet)
        watermark_page = watermark_pdf.pages[0]

        page.merge_page(watermark_page)
        writer.add_page(page)

    with open(output_path, "wb") as f:
        writer.write(f)

    return {"success": True, "output": output_path, "pages_watermarked": len(writer.pages)}

def rotate_pdf(input_path, output_path, angle=90):
    reader = PdfReader(input_path)
    writer = PdfWriter()
    angle = int(angle) % 360

    for page in reader.pages:
        page.rotate(angle)
        writer.add_page(page)

    with open(output_path, "wb") as f:
        writer.write(f)

    return {"success": True, "output": output_path, "rotated_by": angle}

def add_page_numbers(input_path, output_path, format_str="Page {n} of {total}", position="bottom-center"):
    reader = PdfReader(input_path)
    writer = PdfWriter()
    total = len(reader.pages)

    for idx, page in enumerate(reader.pages):
        pw = float(page.mediabox.width)
        ph = float(page.mediabox.height)

        packet = io.BytesIO()
        can = canvas.Canvas(packet, pagesize=(pw, ph))
        can.setFillColor(colors.Color(0.2, 0.2, 0.2, alpha=0.8))
        can.setFont("Helvetica", 10)

        label = format_str.replace("{n}", str(idx + 1)).replace("{total}", str(total))

        if position == "bottom-center":
            can.drawCentredString(pw / 2, 20, label)
        elif position == "bottom-right":
            can.drawRightString(pw - 25, 20, label)
        elif position == "top-center":
            can.drawCentredString(pw / 2, ph - 20, label)

        can.save()
        packet.seek(0)
        num_pdf = PdfReader(packet)
        page.merge_page(num_pdf.pages[0])
        writer.add_page(page)

    with open(output_path, "wb") as f:
        writer.write(f)

    return {"success": True, "output": output_path}

def extract_text(input_path):
    reader = PdfReader(input_path)
    full_text = []
    for idx, page in enumerate(reader.pages):
        full_text.append(f"--- Page {idx + 1} ---\n" + (page.extract_text() or ""))
    return {"success": True, "text": "\n\n".join(full_text), "page_count": len(reader.pages)}

def markdown_to_pdf(markdown_text, output_path):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    
    doc = SimpleDocTemplate(output_path, pagesize=pagesizes.A4)
    styles = getSampleStyleSheet()
    story = []

    lines = markdown_text.split("\n")
    for line in lines:
        line_clean = line.strip()
        if not line_clean:
            story.append(Spacer(1, 10))
        elif line_clean.startswith("# "):
            story.append(Paragraph(line_clean[2:], styles['Title']))
            story.append(Spacer(1, 8))
        elif line_clean.startswith("## "):
            story.append(Paragraph(line_clean[3:], styles['Heading2']))
            story.append(Spacer(1, 6))
        elif line_clean.startswith("### "):
            story.append(Paragraph(line_clean[4:], styles['Heading3']))
            story.append(Spacer(1, 4))
        else:
            story.append(Paragraph(line_clean, styles['Normal']))
            story.append(Spacer(1, 4))

    doc.build(story)
    return {"success": True, "output": output_path}


def detect_script_type(text):
    import re
    if re.search(r'[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿֐-׿]', text):
        return 'rtl-arabic'
    if re.search(r'[ఀ-౿]', text):
        return 'telugu'
    if re.search(r'[ऀ-ॿ]', text):
        return 'devanagari'
    return 'latin'

def pdf_to_word(input_path, output_path):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    reader = PdfReader(input_path)
    doc = docx.Document()

    # Set 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    
    # Title
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(base_name.replace('_', ' ').replace('-', ' ').title())
    title_run.bold = True
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(229, 50, 45) # Brand Red
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(20)

    total_extracted = 0
    for idx, page in enumerate(reader.pages):
        if idx > 0:
            doc.add_page_break()
        text = page.extract_text()
        if text:
            lines = text.split('\n')
            for line in lines:
                line_clean = line.strip()
                if not line_clean:
                    continue
                total_extracted += 1
                st = detect_script_type(line_clean)
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.2
                if st == 'rtl-arabic':
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    run = p.add_run(line_clean)
                    run.font.name = 'Amiri'
                    run.font.size = Pt(13)
                elif st == 'telugu':
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(line_clean)
                    run.font.name = 'Nirmala UI'
                    run.font.size = Pt(12)
                elif st == 'devanagari':
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(line_clean)
                    run.font.name = 'Nirmala UI'
                    run.font.size = Pt(12)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = p.add_run(line_clean)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11)

    # Footer note
    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_p.paragraph_format.space_before = Pt(30)
    footer_run = footer_p.add_run("Converted from PDF to Word by CodeWithAli PDF Tools Suite")
    footer_run.font.size = Pt(9)
    footer_run.font.italic = True
    footer_run.font.color.rgb = RGBColor(128, 128, 128)

    doc.save(output_path)
    return {"success": True, "output": output_path, "pages": len(reader.pages), "lines": total_extracted}

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"error": "No command provided"}))
        sys.exit(1)

    cmd = sys.argv[1]
    data = json.loads(sys.stdin.read()) if not sys.stdin.isatty() else {}

    try:
        if cmd == "merge":
            res = merge_pdfs(data["input_paths"], data["output_path"])
        elif cmd == "split":
            res = split_pdf(data["input_path"], data["output_dir"], data.get("mode", "all"), data.get("page_range", ""))
        elif cmd == "compress":
            res = compress_pdf(data["input_path"], data["output_path"], data.get("level", "recommended"))
        elif cmd == "image_to_pdf":
            res = images_to_pdf(data["image_paths"], data["output_path"], data.get("orientation", "portrait"), data.get("margin", 20), data.get("page_size", "a4"))
        elif cmd == "watermark":
            res = add_watermark(data["input_path"], data["output_path"], data.get("text", "CONFIDENTIAL"), data.get("position", "diagonal"), data.get("opacity", 0.3), data.get("font_size", 48), data.get("color", "#666666"))
        elif cmd == "rotate":
            res = rotate_pdf(data["input_path"], data["output_path"], data.get("angle", 90))
        elif cmd == "page_numbers":
            res = add_page_numbers(data["input_path"], data["output_path"], data.get("format", "Page {n} of {total}"), data.get("position", "bottom-center"))
        elif cmd == "pdf_to_word":
            res = pdf_to_word(data["input_path"], data["output_path"])
        elif cmd == "extract_text":
            res = extract_text(data["input_path"])
        elif cmd == "markdown_to_pdf":
            res = markdown_to_pdf(data["markdown"], data["output_path"])
        else:
            res = {"error": f"Unknown command {cmd}"}
        print(json.dumps(res))
    except Exception as e:
        print(json.dumps({"error": str(e)}))
        sys.exit(1)
